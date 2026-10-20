"""Local-only TLE document/report tools. No server, credential or client-data upload."""
from __future__ import annotations
import json
import re
import subprocess
import sys
import threading
from datetime import datetime, timezone
from pathlib import Path

ROOT = Path(__file__).resolve().parent

def read_pdf(path: Path) -> str:
    from pypdf import PdfReader
    reader = PdfReader(str(path))
    if reader.is_encrypted:
        raise ValueError('Abre y desbloquea el PDF antes de usarlo.')
    if len(reader.pages) > 50:
        raise ValueError('Divide el documento en grupos de hasta 50 páginas.')
    return '\n\n'.join(page.extract_text() or '' for page in reader.pages)

def extract_saved_page(path: Path, source_url: str) -> dict:
    """Scrapling parses a page the owner already saved; it never bypasses a login."""
    from scrapling.parser import Selector
    if path.stat().st_size > 2_000_000:
        raise ValueError('Usa un archivo HTML de menos de 2 MB.')
    if not re.fullmatch(r'https://[^\s<>"\'?#]+', source_url):
        raise ValueError('Indica la URL HTTPS pública original, sin tokens ni parámetros.')
    raw = path.read_text(encoding='utf-8', errors='replace')
    # Ignore code and style before treating any extracted string as page text.
    raw = re.sub(r'<(script|style|noscript|svg|template)\b[^>]*>[\s\S]*?</\1>', ' ', raw, flags=re.I)
    page = Selector(content=raw)
    text = ' '.join(page.css('body *::text').getall())
    return {
        'source_url': source_url,
        'extracted_at': datetime.now(timezone.utc).isoformat(),
        'source_type': 'owner_saved_html_not_live_verified',
        'title': page.css('title::text').get() or '',
        'headings': page.css('h1::text').getall()[:8],
        'public_emails': sorted(set(re.findall(r'[\w.+-]+@[\w.-]+\.[A-Za-z]{2,}', text)))[:20],
        'email_delivery_verified': False,
        'excerpt': re.sub(r'\s+', ' ', text)[:6000],
        'limitations': 'Saved source only. No live freshness, rankings, form or mailbox verification.',
    }

def ocr_document(path: Path, output: Path) -> list:
    from paddleocr import PaddleOCR
    if path.stat().st_size > 30_000_000:
        raise ValueError('Usa documentos de menos de 30 MB.')
    if path.suffix.lower() not in {'.png', '.jpg', '.jpeg', '.webp', '.bmp', '.tif', '.tiff', '.pdf'}:
        raise ValueError('Selecciona una imagen o un PDF.')
    if path.suffix.lower()=='.pdf':
        text=read_pdf(path)
        if len(text.strip())>50:
            output.with_suffix('.txt').write_text(text,encoding='utf-8')
            return [{'mode':'pdf_embedded_text','text_file':str(output.with_suffix('.txt'))}]
    # Native local inference. Initial model downloads require an internet connection;
    # the document itself is never passed to an online OCR endpoint.
    engine=PaddleOCR(lang='en',text_detection_model_name='PP-OCRv6_tiny_det',text_recognition_model_name='PP-OCRv6_tiny_rec',
                    use_doc_orientation_classify=False,use_doc_unwarping=False,use_textline_orientation=False)
    results=[]
    for i,result in enumerate(engine.predict(input=str(path))):
        result.save_to_json(save_path=str(output.parent / (output.stem+'-page-'+str(i+1)+'.json')))
        results.append({'page':i+1,'file':output.stem+'-page-'+str(i+1)+'.json'})
    output.write_text(json.dumps({'source_file':path.name,'mode':'local_ocr_unreviewed','review_required':True,'pages':results},indent=2),encoding='utf-8')
    return results

def make_presentation(source: Path, target: Path) -> None:
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.dml.color import RGBColor
    data=json.loads(source.read_text(encoding='utf-8'))
    if not isinstance(data,dict) or not {'generated_at','active_leads','due','evidence','coverage'} <= data.keys():
        raise ValueError('Usa el JSON exportado desde Sales tools → Reports.')
    deck=Presentation();deck.slide_width=Inches(13.333);deck.slide_height=Inches(7.5)
    def slide(title,content):
        s=deck.slides.add_slide(deck.slide_layouts[6]);s.background.fill.solid();s.background.fill.fore_color.rgb=RGBColor.from_string('FAF8F3')
        def text(value,x,y,w,h,size,color):
            t=s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h)).text_frame;t.word_wrap=True
            t.text=str(value)
            for p in t.paragraphs:p.font.name='Aptos';p.font.size=Pt(size);p.font.color.rgb=RGBColor.from_string(color)
        text('THE LAUNCH ERA / WORKSPACE REPORT',.6,.4,12,.4,11,'205579')
        text(title,.6,1.1,12,1,36,'191919');text(content,.6,2.55,12,3.8,20,'54636B')
        text(str(data['generated_at']),.6,6.9,12,.3,10,'54636B')
    slide('Your sales workspace.',f"Active leads loaded: {data['active_leads']}\nFollow-ups due: {data['due']}\n\nCounts are not revenue, conversion rates or proof of completed work.")
    for offset in range(0,len(data['evidence'][:20]),4):
        refs=data['evidence'][offset:offset+4]
        slide('Source references.', '\n\n'.join(str(r.get('title',''))[:130]+'\n'+str(r.get('source_url') or r.get('id',''))[:180] for r in refs))
    slide('Coverage & limitations.',json.dumps(data['coverage'],indent=2)+'\n\n'+str(data.get('scope','')))
    deck.save(target)

def gui():
    import tkinter as tk
    from tkinter import filedialog, messagebox, simpledialog
    root=tk.Tk();root.title('The Launch Era · Herramientas locales');root.geometry('680x570');root.configure(bg='#FAF8F3')
    status=tk.StringVar(value='Instala las herramientas básicas una vez. No necesitas cuentas ni claves API.')
    tk.Label(root,text='Menos trabajo manual.',font=('Segoe UI',27),bg='#FAF8F3',fg='#191919').pack(padx=25,pady=(28,8),anchor='w')
    tk.Label(root,text='PDFs · OCR local · páginas guardadas · informes editables',bg='#FAF8F3',fg='#205579').pack(padx=25,anchor='w')
    buttons=[]
    def work(fn):
        for b in buttons:b.configure(state='disabled')
        status.set('Trabajando… puedes dejar esta ventana abierta.')
        def runner():
            try:
                result=fn()
                root.after(0,lambda:status.set(str(result or 'Listo. Revisa el archivo generado.')))
            except Exception as e:
                msg=str(e)
                root.after(0,lambda msg=msg:status.set('No se completó: '+msg[:360]))
            finally:
                root.after(0,lambda:[b.configure(state='normal') for b in buttons])
        threading.Thread(target=runner,daemon=True).start()
    def install(name):
        if not messagebox.askyesno('Instalación local','Se descargarán paquetes gratuitos en esta carpeta. OCR también descarga modelos. ¿Continuar?'):return
        def run():
            proc=subprocess.run([sys.executable,'-m','pip','install','--disable-pip-version-check','-r',str(ROOT/name)],capture_output=True,text=True)
            if proc.returncode:raise RuntimeError(proc.stderr[-700:] or 'La instalación no terminó.')
            return 'Instalación terminada. Ya puedes usar esta herramienta.'
        work(run)
    def choose_input(types):
        value=filedialog.askopenfilename(filetypes=types)
        return Path(value) if value else None
    def save(default,ext):
        value=filedialog.asksaveasfilename(initialfile=default,defaultextension=ext)
        return Path(value) if value else None
    def pdf():
        p=choose_input([('PDF','*.pdf')]);out=save('texto-revisar.txt','.txt') if p else None
        if p and out:work(lambda:(out.write_text(read_pdf(p),encoding='utf-8'),'Texto guardado para revisión.')[1])
    def ocr():
        p=choose_input([('Imagen o PDF','*.png *.jpg *.jpeg *.webp *.pdf *.bmp *.tif *.tiff')]);out=save('ocr-revisar.json','.json') if p else None
        if p and out:work(lambda:(ocr_document(p,out),'OCR guardado localmente. Revisa nombres, fechas y cantidades.')[1])
    def research():
        p=choose_input([('HTML guardado','*.html *.htm')]);url=simpledialog.askstring('Fuente','URL pública original (https://…):') if p else None
        out=save('fuente-revisar.json','.json') if url else None
        if p and url and out:work(lambda:(out.write_text(json.dumps(extract_saved_page(p,url),ensure_ascii=False,indent=2),encoding='utf-8'),'Fuente guardada. Los emails no se verificaron.')[1])
    def report():
        p=choose_input([('Informe JSON del Command Center','*.json')]);out=save('TLE-informe.pptx','.pptx') if p else None
        if p and out:work(lambda:(make_presentation(p,out),'PowerPoint editable guardado.')[1])
    for label,fn in [('1. Instalar herramientas básicas',lambda:install('requirements-research.txt')),('2. Instalar OCR local (opcional)',lambda:install('requirements-ocr.txt')),('Extraer texto de un PDF',pdf),('Leer imagen / PDF con OCR',ocr),('Extraer datos de una página HTML guardada (Scrapling)',research),('Convertir informe JSON a PowerPoint',report)]:
        b=tk.Button(root,text=label,command=fn,font=('Segoe UI',11),bg='#BFDEF0',fg='#191919',relief='flat',padx=15,pady=10);b.pack(padx=25,pady=5,fill='x');buttons.append(b)
    tk.Label(root,textvariable=status,bg='#FAF8F3',fg='#405865',wraplength=620,justify='left').pack(padx=25,pady=16,anchor='w')
    root.mainloop()

if __name__=='__main__':
    if not (3,10)<=sys.version_info[:2]<=(3,12):
        raise SystemExit('Usa Python 3.10, 3.11 o 3.12 para compatibilidad con OCR.')
    gui()
